from io import BytesIO

import openpyxl

from app.services.document.base import DocumentExtractor, ExtractedDocument
from app.services.document.errors import DocumentExtractionError, EmptyDocumentError


class ExcelExtractor(DocumentExtractor):
    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        try:
            workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise DocumentExtractionError(
                f"Could not open '{filename}' as an Excel workbook"
            ) from exc

        lines: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))

        raw_text = "\n".join(lines).strip()
        if not raw_text:
            raise EmptyDocumentError(f"'{filename}' contains no data")

        return ExtractedDocument(raw_text=raw_text, source_filename=filename)
